#!/usr/bin/env python3
"""Build a separate study-helps appendix from the current GHSB draft text.

The appendix outline follows the supplied Google Doc, omitting the reading
plan. Verse bodies are rendered from the repository's current OT LXX and NT TR
translation rows rather than from the supplied document's English snippets.
"""

from __future__ import annotations

import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import build_fresh_logos_bible as bible_builder  # noqa: E402
import build_print_proof_bible as print_builder  # noqa: E402
import build_study_bible as study_builder  # noqa: E402


OUT_DIR = ROOT / "output" / "doc"
MD_OUT = OUT_DIR / "greek_heritage_study_helps_appendix.md"
DOCX_OUT = OUT_DIR / "greek_heritage_study_helps_appendix.docx"
DIAG_OUT = OUT_DIR / "greek_heritage_study_helps_appendix_diagnostics.csv"


@dataclass(frozen=True)
class VerseItem:
    ref: str


@dataclass(frozen=True)
class TextItem:
    text: str


@dataclass(frozen=True)
class NumberMeaning:
    number: str
    title: str
    text: str


@dataclass(frozen=True)
class DivineName:
    name: str
    meaning: str
    refs: str


NUMBER_MEANINGS: list[NumberMeaning] = [
    NumberMeaning("1", "Unity, God's Sovereignty", "Represents the oneness of God and the beginning or first place. Key reference: Deuteronomy 6:4."),
    NumberMeaning("2", "Union, Division, or Witness", "Can signify partnership, contrast, or confirmation by two witnesses. Key references: Ecclesiastes 4:9; Matthew 18:16."),
    NumberMeaning("3", "Divine Completeness, Trinity", "Often associated with divine completeness and the triune name. Key references: Matthew 28:19; John 21:14."),
    NumberMeaning("4", "Creation, Universality", "Often associated with the created world and its fullness. Key reference: Revelation 7:1."),
    NumberMeaning("5", "Grace, God's Goodness", "Often tied to God's favor and human weakness met by divine strength. Key reference: Genesis 43:34."),
    NumberMeaning("6", "Humanity, Imperfection", "Associated with humanity and falling short of divine completion. Key references: Genesis 1:26-31; Revelation 13:18."),
    NumberMeaning("7", "Perfection, Completion", "Associated with completed work and fullness. Key references: Genesis 2:2; Revelation's sevenfold imagery."),
    NumberMeaning("8", "New Beginnings", "Associated with renewal and a new order. Key references: Genesis 7:13; Leviticus 12:3."),
    NumberMeaning("9", "Divine Completeness, Finality", "Often used as a number of completeness or finality. Key reference: Galatians 5:22-23."),
    NumberMeaning("10", "Divine Order, Completeness of Order", "Often marks a complete ordered set or human responsibility. Key references: Exodus 20; Exodus 7-12."),
    NumberMeaning("11", "Disorder, Chaos, Judgment", "Sometimes associated with disorder or judgment."),
    NumberMeaning("12", "God's Government, Authority", "Associated with divine order and covenant administration. Key references: Genesis 49:28; Matthew 10:1-2."),
    NumberMeaning("13", "Rebellion, Lawlessness", "Sometimes associated with rebellion and defilement."),
    NumberMeaning("14", "Double Spiritual Perfection", "A doubled seven. Matthew structures Jesus' genealogy in groups of fourteen."),
    NumberMeaning("15", "Rest After Deliverance", "Often associated with rest following deliverance. Key references: Genesis 15:12-16; Exodus 12:40-42."),
    NumberMeaning("16", "Love and Loving", "Often associated with love. Key references: Genesis 46:18; 1 Corinthians 13:4-8."),
    NumberMeaning("17", "Overcoming the Enemy, Victory", "Often treated as ten plus seven, joining order and completion."),
    NumberMeaning("18", "Bondage", "Associated with bondage or oppression. Key references: Judges 3:14; Luke 13:11-12."),
    NumberMeaning("19", "God's Perfect Order in Judgment", "Often associated with ordered judgment. Key references: Joshua 19:38; Psalm 19."),
    NumberMeaning("20", "Perfect Waiting", "Often associated with waiting, maturity, or preparation."),
    NumberMeaning("21", "Great Wickedness, Rebellion, and Sin", "Associated with the full outgrowth of sin and rebellion. Key references: 2 Timothy 3:1-5; Revelation 21:21."),
    NumberMeaning("22", "Light", "Associated with light and revelation. Key references: Psalm 119:105; John 12:46."),
    NumberMeaning("23", "Divine Protection and Guidance", "Associated with protection and guidance. Key references: Psalm 23; 1 Corinthians 10:8."),
    NumberMeaning("24", "Priesthood", "Associated with priestly service and worship. Key references: 1 Chronicles 23-24; Revelation 4:4."),
    NumberMeaning("25", "Grace Upon Grace", "Associated with multiplied grace. Key references: John 1:16-17; Ezekiel 40:13."),
    NumberMeaning("26", "Eternal God", "Often associated with the covenant name by Hebrew reckoning. Key reference: Genesis 1:26."),
    NumberMeaning("30", "Dedication to a Task", "Associated with maturity, public service, and dedication."),
    NumberMeaning("33", "God's Promises", "Often connected with promise and completion."),
    NumberMeaning("40", "Testing, Trial, Preparation", "Associated with testing and preparation. Key references: Genesis 7:12; Matthew 4:2."),
    NumberMeaning("50", "Holy Spirit, Pentecost, Deliverance", "Associated with Pentecost, jubilee, and deliverance."),
    NumberMeaning("60", "Threshold of Later Life, Pride", "Sometimes associated with later life or human strength."),
    NumberMeaning("70", "Completeness, Restoration", "Associated with fullness, restoration, and appointed periods. Key references: Jeremiah 25:11; Luke 10:1."),
    NumberMeaning("666", "Perfection of Imperfection, Beast", "Associated with the beast and ultimate human opposition to God. Key reference: Revelation 13:18."),
]


DIVINE_NAMES: list[DivineName] = [
    DivineName("El Shaddai", "God Almighty", "Genesis 17:1"),
    DivineName("El Elyon", "Most High God", "Daniel 3:26; Psalm 7:17"),
    DivineName("Adonai", "Lord; Master", "Deuteronomy 3:24; Psalm 16:2"),
    DivineName("Yahweh", "The Lord; covenant name", "Genesis 2:4; Exodus 3:14-15"),
    DivineName("Yahweh Nissi", "The Lord my banner", "Exodus 17:15"),
    DivineName("Yahweh Raah", "The Lord my shepherd", "Psalm 23:1"),
    DivineName("Yahweh Rapha", "The Lord who heals", "Exodus 15:26"),
    DivineName("Yahweh Shammah", "The Lord is there", "Ezekiel 48:35"),
    DivineName("Yahweh Tsidkenu", "The Lord our righteousness", "Jeremiah 23:6"),
    DivineName("Yahweh Mekoddishkem", "The Lord who sanctifies you", "Exodus 31:13"),
    DivineName("El Olam", "Everlasting God", "Genesis 21:33; Psalm 48:14"),
    DivineName("Elohim", "God", "Genesis 1:1"),
    DivineName("Qanna", "Jealous", "Exodus 20:5; Exodus 34:14"),
    DivineName("Yahweh Jireh", "The Lord will provide", "Genesis 22:14"),
    DivineName("Yahweh Shalom", "The Lord is peace", "Judges 6:24"),
    DivineName("Yahweh Sabaoth", "Lord of hosts", "1 Samuel 17:45"),
]


VERSE_SECTIONS: list[tuple[str, list[VerseItem | TextItem]]] = [
    ("Faith and Trust in God", [VerseItem("Proverbs 3:5-6"), VerseItem("Hebrews 11:1"), VerseItem("2 Corinthians 5:7"), VerseItem("Philippians 4:13"), VerseItem("Romans 8:28"), VerseItem("Matthew 17:20"), VerseItem("Psalm 37:5")]),
    ("God's Love and Salvation", [VerseItem("John 3:16"), VerseItem("Romans 5:8"), VerseItem("Ephesians 2:8-9"), VerseItem("John 14:6"), VerseItem("Romans 10:9"), VerseItem("Romans 8:38-39"), VerseItem("Titus 3:5")]),
    ("Strength and Comfort", [VerseItem("Psalm 23:1"), VerseItem("Isaiah 41:10"), VerseItem("Psalm 46:1"), VerseItem("Matthew 11:28"), VerseItem("1 Peter 5:7"), VerseItem("Ephesians 6:11-17"), VerseItem("Philippians 4:19"), VerseItem("2 Corinthians 12:9")]),
    ("Peace and Hope", [VerseItem("Jeremiah 29:11"), VerseItem("Romans 15:13"), VerseItem("Philippians 4:6-7"), VerseItem("John 16:33"), VerseItem("Psalm 34:18"), VerseItem("Isaiah 26:3"), VerseItem("Romans 5:1")]),
    ("Guidance and Wisdom", [VerseItem("James 1:5"), VerseItem("Psalm 119:105"), VerseItem("Micah 6:8"), VerseItem("Matthew 6:33"), VerseItem("Joshua 1:9"), VerseItem("Proverbs 16:9"), VerseItem("Isaiah 30:21")]),
    ("Love and Relationships", [VerseItem("1 Corinthians 13:4-7"), VerseItem("1 John 4:19"), VerseItem("Colossians 3:13"), VerseItem("Ephesians 4:32"), VerseItem("Mark 12:30-31"), VerseItem("Ephesians 5:25"), VerseItem("1 Peter 4:8")]),
    ("Encouragement in Trials", [VerseItem("James 1:2-3"), VerseItem("Romans 12:12"), VerseItem("1 Corinthians 10:13"), VerseItem("Isaiah 40:31"), VerseItem("Psalm 55:22"), VerseItem("2 Timothy 1:7"), VerseItem("Hebrews 12:1")]),
    ("Obedience and Righteous Living", [VerseItem("Romans 12:2"), VerseItem("Galatians 5:22-23"), VerseItem("Matthew 5:16"), VerseItem("1 Thessalonians 5:16-18"), VerseItem("Psalm 1:1-2"), VerseItem("John 14:15"), VerseItem("Proverbs 4:23")]),
    ("God's Power and Glory", [VerseItem("Genesis 1:1"), VerseItem("Revelation 21:4"), VerseItem("Psalm 19:1"), VerseItem("Colossians 3:23"), VerseItem("Matthew 28:19-20"), VerseItem("Ephesians 3:20"), VerseItem("Psalm 139:14")]),
    ("Repentance and Forgiveness", [VerseItem("1 John 1:9"), VerseItem("Acts 3:19"), VerseItem("Psalm 51:10"), VerseItem("2 Chronicles 7:14"), VerseItem("Romans 6:23"), VerseItem("Luke 15:7"), VerseItem("Matthew 6:14-15")]),
    ("The Return of Christ", [VerseItem("Matthew 24:30-31"), VerseItem("Acts 1:10-11"), VerseItem("Revelation 1:7"), VerseItem("1 Thessalonians 4:16-17"), VerseItem("Titus 2:13"), VerseItem("Matthew 25:31"), VerseItem("Revelation 22:12")]),
    ("Signs of the End", [VerseItem("Matthew 24:3-8"), VerseItem("Luke 21:25-28"), VerseItem("2 Timothy 3:1-5"), VerseItem("Daniel 12:4"), VerseItem("Matthew 24:12"), VerseItem("Mark 13:7-8"), VerseItem("Romans 8:22")]),
    ("The Rise of Deception and Antichrist", [VerseItem("Matthew 24:24"), VerseItem("2 Thessalonians 2:3-4"), VerseItem("1 John 2:18"), VerseItem("Revelation 13:4-8"), VerseItem("Daniel 7:25"), VerseItem("Revelation 13:11-14"), VerseItem("1 Timothy 4:1")]),
    ("The Great Tribulation", [VerseItem("Matthew 24:21-22"), VerseItem("Daniel 9:27"), VerseItem("Revelation 6:3-8"), VerseItem("Revelation 7:14"), VerseItem("Daniel 12:1"), VerseItem("Revelation 13:7"), VerseItem("Revelation 16:1")]),
    ("Israel and the Nations", [VerseItem("Zechariah 12:2-3"), VerseItem("Ezekiel 38:8-9"), VerseItem("Romans 11:25-26"), VerseItem("Matthew 24:14"), VerseItem("Joel 3:2"), VerseItem("Isaiah 11:11-12"), VerseItem("Amos 9:14-15")]),
    ("The Day of the Lord and Judgment", [VerseItem("Joel 2:31"), VerseItem("1 Thessalonians 5:2-3"), VerseItem("2 Peter 3:10"), VerseItem("Zephaniah 1:14-15"), VerseItem("Revelation 6:15-17"), VerseItem("Revelation 20:11-12"), VerseItem("Malachi 4:1")]),
    ("The Millennium and Eternal Kingdom", [VerseItem("Revelation 20:1-4"), VerseItem("Isaiah 2:2-4"), VerseItem("Isaiah 11:6-9"), VerseItem("Revelation 21:1-4"), VerseItem("2 Peter 3:13"), VerseItem("Isaiah 65:17"), VerseItem("Revelation 22:3-5")]),
    ("Watchfulness and Readiness", [VerseItem("Matthew 24:42-44"), VerseItem("Matthew 25:13"), VerseItem("Luke 21:36"), VerseItem("Revelation 3:11"), VerseItem("1 Peter 4:7"), VerseItem("1 Thessalonians 5:6"), VerseItem("Mark 13:35-37")]),
    (
        "The Day of the Lord and the Day of Christ",
        [
            TextItem("The Day of the Lord: a time of divine judgment and upheaval on the earth, focused on God's justice against sin and the ungodly."),
            TextItem("The Day of Christ: a time of completion, reward, and manifestation for believers, focused on Christ's appearing and his people."),
            TextItem("Day of the Lord"),
            VerseItem("Isaiah 13:6"),
            VerseItem("Joel 2:1-2"),
            VerseItem("Amos 5:18-20"),
            VerseItem("Zephaniah 1:14-15"),
            VerseItem("1 Thessalonians 5:2-3"),
            VerseItem("2 Peter 3:10"),
            VerseItem("Revelation 6:17"),
            TextItem("Day of Christ"),
            VerseItem("Philippians 1:6"),
            VerseItem("Philippians 1:10"),
            VerseItem("Philippians 2:16"),
            VerseItem("1 Corinthians 1:8"),
            VerseItem("2 Corinthians 1:14"),
            VerseItem("2 Thessalonians 2:2"),
            VerseItem("Romans 14:10-12"),
        ],
    ),
    (
        "Kingdom of Heaven and Kingdom of God",
        [
            TextItem("Kingdom of Heaven: Matthew's usual phrase for the heavenly reign brought near through Messiah."),
            TextItem("Kingdom of God: the broader phrase for God's reign, entered by new birth and faith in Christ."),
            TextItem("Kingdom of Heaven"),
            VerseItem("Matthew 3:2"),
            VerseItem("Matthew 5:3"),
            VerseItem("Matthew 13:24"),
            VerseItem("Matthew 13:47"),
            VerseItem("Matthew 18:3"),
            VerseItem("Matthew 22:2"),
            VerseItem("Matthew 25:1"),
            TextItem("Kingdom of God"),
            VerseItem("Mark 1:14-15"),
            VerseItem("Luke 17:20-21"),
            VerseItem("John 3:3"),
            VerseItem("Romans 14:17"),
            VerseItem("1 Corinthians 6:9-10"),
            VerseItem("Colossians 1:13"),
            VerseItem("2 Peter 1:11"),
        ],
    ),
]


def load_verses() -> dict[str, dict[str, str]]:
    verses: dict[str, dict[str, str]] = {}
    for path in [ROOT / "data/raw/lxx_greek/ot_full.csv", ROOT / "data/raw/tr_greek/nt_full.csv"]:
        with path.open(newline="", encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                code_ref = f"{row['book_code']} {row['chapter']}:{row['verse']}"
                verses[code_ref] = row
    return verses


VERSE_INDEX = load_verses()
ENGLISH_TO_LOCAL = bible_builder.invert_versification_map(
    bible_builder.effective_versification_map(None)
)
REF_BOOK_LABEL_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])("
    + "|".join(re.escape(label) for label in sorted(study_builder.REF_BOOK_ALIASES, key=len, reverse=True))
    + r")\s+\d+(?::\d+(?:-\d+)?)?(?:-\d+(?::\d+)?)?"
)


def mapped_code_ref(code: str, chapter: int, verse: int) -> str:
    if code == "DAG":
        code = "DAN"
    standard_ref = f"{code} {chapter}:{verse}"
    if code == "DAN":
        return standard_ref
    return ENGLISH_TO_LOCAL.get(standard_ref, standard_ref)


def normalize_ref_text(ref: str) -> str:
    return (
        ref.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u00a0", " ")
        .replace("Psalm ", "Psalms ")
        .strip()
    )


def abbreviate_ref_label(ref: str) -> str:
    return print_builder.abbreviate_print_crossref(ref.replace("Psalm ", "Psalms "))


def abbreviate_embedded_references(text: str) -> str:
    return REF_BOOK_LABEL_PATTERN.sub(lambda match: abbreviate_ref_label(match.group(0)), text)


def local_ref_label(original_ref: str, code_refs: list[str]) -> str:
    local_labels = [bible_builder.display_ref_from_code_ref(ref) or ref for ref in code_refs]
    if not local_labels:
        return original_ref
    parsed_labels = [study_builder.parse_cross_reference(label) for label in local_labels]
    if all(parsed_labels) and len(parsed_labels) > 1:
        first = parsed_labels[0]
        last = parsed_labels[-1]
        if (
            first
            and last
            and first[0] == last[0]
            and first[2] == last[2]
            and first[3] is not None
            and last[3] is not None
            and all((p and p[0] == first[0] and p[2] == first[2]) for p in parsed_labels)
        ):
            local_label = f"{first[1]} {first[2]}:{first[3]}-{last[3]}"
        else:
            local_label = "; ".join(local_labels)
    else:
        local_label = local_labels[0]
    if normalize_ref_text(local_label) != normalize_ref_text(original_ref):
        return f"{abbreviate_ref_label(local_label)} (English {abbreviate_ref_label(original_ref)})"
    return abbreviate_ref_label(local_label)


def expand_reference(ref: str) -> tuple[str, str, list[str]]:
    parsed = study_builder.parse_cross_reference(ref)
    if not parsed:
        return ref, f"[unresolved reference: {ref}]", [ref]
    code, _book_label, chapter, verse_start, verse_end, raw_ref = parsed
    if verse_start is None:
        return raw_ref, f"[chapter reference retained: {raw_ref}]", [raw_ref]
    verse_end = verse_end or verse_start
    code_refs = [mapped_code_ref(code, chapter, verse) for verse in range(verse_start, verse_end + 1)]
    parts: list[str] = []
    unresolved: list[str] = []
    for code_ref in code_refs:
        row = VERSE_INDEX.get(code_ref)
        if not row:
            unresolved.append(code_ref)
            continue
        text = bible_builder.normalize_bible_text_for_output(row["draft_translation"]).strip()
        local_verse = code_ref.rsplit(":", 1)[1]
        if len(code_refs) == 1:
            parts.append(text)
        else:
            parts.append(f"{local_verse} {text}")
    label = local_ref_label(raw_ref, code_refs)
    if unresolved:
        return label, f"[missing local verse text: {', '.join(unresolved)}]", unresolved
    return label, " ".join(parts), []


def normalize_reference_list(refs: str) -> str:
    normalized: list[str] = []
    for piece in re.split(r";\s*", refs):
        piece = piece.strip()
        if not piece:
            continue
        parsed = study_builder.parse_cross_reference(piece)
        if parsed and parsed[3] is not None:
            code, _book, chapter, verse_start, verse_end, raw_ref = parsed
            code_refs = [mapped_code_ref(code, chapter, verse_start)]
            if verse_end and verse_end != verse_start:
                code_refs.append(mapped_code_ref(code, chapter, verse_end))
            normalized.append(local_ref_label(raw_ref, code_refs))
        else:
            normalized.append(abbreviate_ref_label(piece))
    return "; ".join(abbreviate_ref_label(item) for item in normalized)


def render_markdown() -> tuple[str, list[tuple[str, str]]]:
    diagnostics: list[tuple[str, str]] = []
    lines: list[str] = [
        "# Study Helps Appendix",
        "",
        "This separate appendix draft is adapted from the supplied insert with the reading plan omitted. Scripture lines are rendered from the current Greek Heritage Study Bible draft text.",
        "",
        "## Number Meanings",
        "",
    ]
    for item in NUMBER_MEANINGS:
        lines.append(f"**{item.number}. {item.title}.** {abbreviate_embedded_references(item.text)}")
        lines.append("")
    lines.extend(
        [
            "## Divine Names and Titles",
            "",
            "| Name | Meaning | References |",
            "|---|---|---|",
        ]
    )
    for item in DIVINE_NAMES:
        lines.append(f"| {item.name} | {item.meaning} | {normalize_reference_list(item.refs)} |")
    lines.append("")
    lines.append("## Thematic Passages")
    lines.append("")
    for heading, items in VERSE_SECTIONS:
        lines.append(f"### {heading}")
        lines.append("")
        for item in items:
            if isinstance(item, TextItem):
                lines.append(item.text)
                lines.append("")
                continue
            label, text, unresolved = expand_reference(item.ref)
            if unresolved:
                diagnostics.append((item.ref, "; ".join(unresolved)))
            lines.append(f"**{label}.** {text}")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n", diagnostics


def render_docx(markdown: str) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_BREAK
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover - environment fallback
        raise SystemExit(f"python-docx is required for DOCX output: {exc}") from exc

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in [("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)]:
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("| Name | Meaning | References |"):
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            headers = ["Name", "Meaning", "References"]
            for cell, header in zip(table.rows[0].cells, headers):
                cell.text = header
            index += 2
            while index < len(lines) and lines[index].startswith("| "):
                cells = [part.strip() for part in lines[index].strip("|").split("|")]
                row = table.add_row().cells
                for cell, value in zip(row, cells):
                    cell.text = value
                index += 1
            continue
        elif line.startswith("**") and ".** " in line:
            marker = line.find(".** ")
            title = line[2 : marker + 1]
            body = line[marker + 4 :]
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(title + " ")
            run.bold = True
            paragraph.add_run(body)
        elif line:
            doc.add_paragraph(line)
        index += 1
    doc.save(DOCX_OUT)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    markdown, diagnostics = render_markdown()
    MD_OUT.write_text(markdown, encoding="utf-8")
    with DIAG_OUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["input_ref", "issue"])
        writer.writerows(diagnostics)
    render_docx(markdown)
    print(f"wrote {MD_OUT}")
    print(f"wrote {DOCX_OUT}")
    print(f"wrote {DIAG_OUT}")
    print(f"unresolved={len(diagnostics)}")


if __name__ == "__main__":
    main()
